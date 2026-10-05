import subprocess
import openpyxl
import zipfile
import os

filepath = r"C:\Users\HP\Downloads\Test_Suite_US-101_US-103.xlsx"
print(f"File exists: {os.path.exists(filepath)}, size: {os.path.getsize(filepath)}")

# 1. Test zipfile validity
try:
    with zipfile.ZipFile(filepath, 'r') as z:
        print("Zip entries:", len(z.namelist()))
        for name in z.namelist()[:10]:
            print("  ", name)
        # Read xl/worksheets/sheet1.xml
        for s in ['xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml', 'xl/worksheets/sheet3.xml', 'xl/workbook.xml']:
            if s in z.namelist():
                content = z.read(s).decode('utf-8', errors='replace')
                print(f"=== {s} (first 300 chars) ===")
                print(content[:300])
except Exception as e:
    print("Zipfile error:", e)

# 2. Test openpyxl
try:
    wb = openpyxl.load_workbook(filepath)
    print("openpyxl loaded sheetnames:", wb.sheetnames)
    for name in wb.sheetnames:
        sheet = wb[name]
        print(f"Sheet {name}: max_row={sheet.max_row}, max_col={sheet.max_column}")
except Exception as e:
    print("openpyxl load error:", e)

# 3. Test Excel COM
ps_script = f'''
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
try {{
    $wb = $excel.Workbooks.Open("{filepath}")
    Write-Host "Excel COM opened file successfully!"
    $wb.Close($false)
}} catch {{
    Write-Host "Excel COM Exception: " $($_.Exception.Message)
}} finally {{
    $excel.Quit()
    [System.Runtime.Interopservices.Marshal]::ReleaseComObject($excel) | Out-Null
}}
'''
with open("test_open_user_file.ps1", "w") as f:
    f.write(ps_script)

res = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", "test_open_user_file.ps1"], capture_output=True, text=True)
print("PowerShell STDOUT:", res.stdout)
print("PowerShell STDERR:", res.stderr)
