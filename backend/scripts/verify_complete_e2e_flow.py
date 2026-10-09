"""
End-to-End Workflow Verification Script for Azure DevOps Local Mock Integration.
Validates all 16 steps specified in Requirement 22.
"""
import sys
import os
import io
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from fastapi.testclient import TestClient
import openpyxl
from app.main import app
from app.core.config import settings
from app.services.azure_devops import get_azure_devops_service
from app.db.session import SessionLocal
from app.models.test_case import TestCaseRecord

client = TestClient(app)
AUTH_HEADERS = {"Authorization": "Bearer test-token-admin"}


def run_e2e_verification():
    print("=" * 70)
    print("AZURE DEVOPS LOCAL MOCK SERVICE - END-TO-END FLOW VERIFICATION")
    print("=" * 70)

    # Clean up any pre-existing TC-1001 test cases in the test database for clean assertions
    db = SessionLocal()
    try:
        db.query(TestCaseRecord).filter(TestCaseRecord.id.like("TC-1001-%")).delete()
        db.commit()
    finally:
        db.close()

    # Step 1 & 2: Backend and Frontend environment configuration
    print("\n[Step 1 & 2] Verifying Service Configuration...")
    settings.AZURE_DEVOPS_MODE = "mock"
    ado_service = get_azure_devops_service()
    assert ado_service.mode == "mock", f"Expected mock mode, got {ado_service.mode}"
    print(f"  [OK] Service mode: {ado_service.mode} (Local Mock Engine active)")

    # Step 3: Connect to mock Azure DevOps
    print("\n[Step 3] Connecting to Mock Azure DevOps...")
    connect_payload = {
        "organization": "DemoOrganization",
        "project": "TestCaseDemoProject"
    }
    resp = client.post("/api/v1/devops/azure/connect", json=connect_payload, headers=AUTH_HEADERS)
    assert resp.status_code == 200, f"Connect failed: {resp.text}"
    conn_data = resp.json()
    assert conn_data["connected"] is True
    assert conn_data["organization"] == "DemoOrganization"
    assert conn_data["project"] == "TestCaseDemoProject"
    assert conn_data["mode"] == "mock"
    print(f"  [OK] Connected: {conn_data['connected']} | Org: {conn_data['organization']} | Project: {conn_data['project']}")

    # Step 4: Select Project
    print("\n[Step 4] Retrieving Available Mock Projects...")
    resp = client.get("/api/v1/devops/azure/projects", headers=AUTH_HEADERS)
    assert resp.status_code == 200
    projects = resp.json()
    print(f"  [OK] Available Projects ({len(projects)} found):")
    for p in projects:
        print(f"    - [{p['id']}] {p['name']} ({p['state']})")
    assert any(p["name"] == "TestCaseDemoProject" for p in projects)

    # Step 5: Select Area
    print("\n[Step 5] Selecting Area Hierarchy from Mock Azure DevOps...")
    resp = client.get("/api/v1/devops/azure/areas", headers=AUTH_HEADERS)
    assert resp.status_code == 200
    areas = resp.json()
    print(f"  [OK] Retrieved Areas ({len(areas)} found):")
    for a in areas:
        print(f"    - ID: {a['id']} | Name: {a['name']:<20} | Path: {a['path']}")
    auth_area = next((a for a in areas if a["name"] == "Authentication"), None)
    assert auth_area is not None, "Authentication area must be present"

    # Step 6: Retrieve User Stories
    print(f"\n[Step 6] Retrieving User Stories for Area '{auth_area['name']}'...")
    resp = client.get(
        f"/api/v1/devops/azure/work-items?area_path={auth_area['path']}&work_item_type=User Story",
        headers=AUTH_HEADERS
    )
    assert resp.status_code == 200
    stories = resp.json()
    assert len(stories) >= 1
    print(f"  [OK] Found {len(stories)} User Stories:")
    for s in stories:
        print(f"    - #{s['id']}: {s['title']} [State: {s['state']}]")

    # Step 7: Select a Story
    selected_story_id = 1001
    print(f"\n[Step 7] Selecting User Story #{selected_story_id}...")

    # Step 8: View Story Details
    print(f"\n[Step 8] Viewing User Story Details for #{selected_story_id}...")
    resp = client.get(f"/api/v1/devops/azure/work-items/{selected_story_id}?project=TestCaseDemoProject", headers=AUTH_HEADERS)
    assert resp.status_code == 200
    story_details = resp.json()
    assert int(story_details["id"]) == 1001
    assert "User Login" in story_details["title"]
    print(f"  [OK] ID: {story_details['id']}")
    print(f"  [OK] Title: {story_details['title']}")
    print(f"  [OK] Area Path: {story_details.get('area_path')}")
    print(f"  [OK] Description: {story_details['description']}")
    print(f"  [OK] Acceptance Criteria:\n    " + story_details['acceptance_criteria'].replace("\n", "\n    "))

    # Step 9: Generate Test Cases using Existing Engine
    print(f"\n[Step 9] Generating Test Cases for User Story #{selected_story_id}...")
    gen_payload = {
        "application_id": 1,
        "requirement_source": "USER_STORY",
        "user_story_ids": ["1001"],
        "requirement_identifiers": ["1001"]
    }
    resp = client.post("/api/v1/generation/create", json=gen_payload, headers=AUTH_HEADERS)
    assert resp.status_code == 200, f"Generation failed: {resp.text}"
    gen_data = resp.json()

    pos_cases = gen_data["positive_test_cases"]
    neg_cases = gen_data["negative_test_cases"]
    edge_cases = gen_data["edge_test_cases"]
    all_cases = pos_cases + neg_cases + edge_cases

    assert len(pos_cases) > 0, "Expected positive test cases"
    assert len(neg_cases) > 0, "Expected negative test cases"
    assert len(edge_cases) > 0, "Expected edge test cases"
    print(f"  [OK] Generated {len(all_cases)} test cases:")
    print(f"    - Positive: {len(pos_cases)}")
    print(f"    - Negative: {len(neg_cases)}")
    print(f"    - Edge:     {len(edge_cases)}")
    for tc in all_cases[:3]:
        print(f"    Sample: [{tc['id']}] ({tc.get('category')}) {tc.get('scenario')} -> Work Item #{tc.get('azure_work_item_id')}")
        assert tc["id"].startswith(f"TC-{selected_story_id}-")

    # Step 10: Review / Edit Test Cases
    print("\n[Step 10] Reviewing and Editing Test Case (simulating user edit in UI)...")
    tc_to_edit = all_cases[0]
    tc_to_edit["scenario"] = tc_to_edit["scenario"] + " [Reviewed & Edited]"
    tc_to_edit["expected_result"] = tc_to_edit["expected_result"] + " (Verified per security guidelines)"
    print(f"  [OK] Updated scenario: {tc_to_edit['scenario']}")
    print(f"  [OK] Updated expected_result: {tc_to_edit['expected_result']}")

    # Step 11: Download Excel
    print("\n[Step 11] Downloading Excel Export...")
    export_payload = {
        "positive_test_cases": pos_cases,
        "negative_test_cases": neg_cases,
        "edge_test_cases": edge_cases,
        "application_name": "Authentication",
        "requirement_ref": "1001"
    }
    resp = client.post("/api/v1/export/excel", json=export_payload, headers=AUTH_HEADERS)
    assert resp.status_code == 200
    assert resp.headers["content-type"] in [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream"
    ]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    print(f"  [OK] Excel workbook generated successfully. Sheet names: {wb.sheetnames}")
    assert set(wb.sheetnames) == {"Positive", "Negative", "Edge"}

    # Step 12: Sync Test Cases to Azure DevOps
    print("\n[Step 12] Syncing Test Cases to Azure DevOps Mock Service...")
    sync_payload = {
        "test_cases": [tc_to_edit],
        "work_item_id": str(selected_story_id),
        "project": "TestCaseDemoProject"
    }
    resp = client.post("/api/v1/devops/azure/sync", json=sync_payload, headers=AUTH_HEADERS)
    assert resp.status_code == 200, f"Sync failed: {resp.text}"
    sync_cases = resp.json()["test_cases"]
    first_result = sync_cases[0]
    print(f"  [OK] First Sync Result:")
    print(f"    - Test Case: {first_result['id']}")
    print(f"    - Azure ADO ID: {first_result['azure_test_case_id']}")
    print(f"    - Status: {first_result['sync_status']}")

    # Step 13: Verify mock ADO IDs are returned
    print("\n[Step 13] Verifying Mock ADO ID format and persistence...")
    assert first_result["azure_test_case_id"] is not None
    assert first_result["azure_test_case_id"] >= 50000
    assert first_result["sync_status"] == "Synced"
    assigned_ado_id = first_result["azure_test_case_id"]
    print(f"  [OK] Assigned ADO ID: {assigned_ado_id} (Starting range >= 50000)")

    # Step 14: Sync again
    print(f"\n[Step 14] Triggering Second Sync with already synced ADO ID ({assigned_ado_id})...")
    # Simulate user keeping the azure_test_case_id on subsequent review/sync
    tc_to_edit["azure_test_case_id"] = assigned_ado_id
    tc_to_edit["expected_result"] = "Second revision expected result"
    sync_again_payload = {
        "test_cases": [tc_to_edit],
        "work_item_id": str(selected_story_id),
        "project": "TestCaseDemoProject"
    }
    resp_again = client.post("/api/v1/devops/azure/sync", json=sync_again_payload, headers=AUTH_HEADERS)
    assert resp_again.status_code == 200
    second_sync_cases = resp_again.json()["test_cases"]
    second_result = second_sync_cases[0]

    # Step 15 & 16: Verify duplicates are NOT created, and existing test case is updated
    print("\n[Step 15 & 16] Verifying Duplicate Prevention and Update Behavior...")
    print(f"  [OK] Second Sync Result:")
    print(f"    - Test Case: {second_result['id']}")
    print(f"    - Azure ADO ID: {second_result['azure_test_case_id']}")
    print(f"    - Status: {second_result['sync_status']}")
    assert second_result["azure_test_case_id"] == assigned_ado_id, (
        f"Duplicate created! Expected {assigned_ado_id}, got {second_result['azure_test_case_id']}"
    )
    assert second_result["sync_status"] == "Updated", f"Expected 'Updated', got {second_result['sync_status']}"
    print(f"  [OK] Confirmed: ADO ID {assigned_ado_id} preserved without duplicate creation!")
    print(f"  [OK] Confirmed: Status marked as '{second_result['sync_status']}'")

    print("\n" + "=" * 70)
    print("ALL 16 WORKFLOW STEPS VERIFIED SUCCESSFULLY AND PASSED!")
    print("=" * 70)


if __name__ == "__main__":
    run_e2e_verification()
