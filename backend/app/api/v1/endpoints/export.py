import re
import io
import logging
import openpyxl
from fastapi import APIRouter, Depends, Response, HTTPException
from app.schemas.test_case import ExportExcelRequest
from app.services.excel_exporter import excel_exporter
from app.models.user import User
from app.api.deps import get_current_user

logger = logging.getLogger(__name__)

router = APIRouter()

def _build_excel_file(payload: ExportExcelRequest):
    excel_stream = excel_exporter.generate_workbook(
        positive_cases=payload.positive_test_cases,
        negative_cases=payload.negative_test_cases,
        edge_cases=payload.edge_test_cases,
        app_name=payload.application_name,
        req_ref=payload.requirement_ref
    )

    all_cases = (payload.positive_test_cases or []) + (payload.negative_test_cases or []) + (payload.edge_test_cases or [])
    story_ids = []
    for tc in all_cases:
        us = getattr(tc, 'user_story', None)
        if us and str(us).strip() and str(us).strip() not in story_ids:
            story_ids.append(str(us).strip())

    if story_ids:
        clean_tokens = [re.sub(r'[^a-zA-Z0-9_\-]', '_', s.strip()) for s in story_ids if s.strip()]
    else:
        clean_tokens = [re.sub(r'[^a-zA-Z0-9_\-]', '_', s.strip()) for s in re.split(r'[,;\s]+', payload.requirement_ref or '') if s.strip()]

    clean_tokens = [t for t in clean_tokens if t]
    clean_ref = "_".join(clean_tokens[:3]) if clean_tokens else "US-101"
    clean_ref = re.sub(r'_+', '_', clean_ref).strip('_')[:50]
    if not clean_ref:
        clean_ref = "US-101"

    # Step 3 & Step 5 requirement: The filename MUST ALWAYS end with .xlsx (never .xls)
    clean_ref = re.sub(r'\.(xlsx|xls)$', '', clean_ref, flags=re.IGNORECASE)
    filename = f"Test_Suite_{clean_ref}.xlsx"
    content = excel_stream.getvalue()

    # Step 2 requirement: Validate ZIP / OOXML binary file signature (PK\x03\x04)
    if len(content) < 4 or content[:4] != b'PK\x03\x04':
        raise ValueError("Generated Excel content has invalid OOXML ZIP binary signature")

    # Step 2 requirement: Validate openpyxl workbook can be loaded from bytes
    validation_wb = openpyxl.load_workbook(io.BytesIO(content))
    expected_sheets = {"Positive", "Negative", "Edge"}
    found_sheets = set(validation_wb.sheetnames)
    if not expected_sheets.issubset(found_sheets):
        validation_wb.close()
        raise ValueError(f"Workbook missing required sheets. Found: {validation_wb.sheetnames}")
    validation_wb.close()

    logger.info(
        f"[EXPORT_ENDPOINT] Built valid XLSX export:\n"
        f" - Filename: {filename}\n"
        f" - File size: {len(content)} bytes\n"
        f" - Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\n"
        f" - Binary Signature: {content[:4]!r}\n"
        f" - Sheets: {list(found_sheets)}"
    )

    return content, filename


@router.post("/excel")
async def export_excel(payload: ExportExcelRequest, current_user: User = Depends(get_current_user)):
    """
    Canonical binary POST export: Returns valid .xlsx binary stream.
    Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet
    Content-Disposition: attachment; filename="Test_Suite_...xlsx"
    """
    try:
        content, filename = _build_excel_file(payload)
        return Response(
            content=content,
            status_code=200,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "Content-Disposition": f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}',
                "Content-Length": str(len(content)),
                "Access-Control-Expose-Headers": "Content-Disposition, Content-Length",
                "Cache-Control": "no-cache, no-store, must-revalidate",
                "Pragma": "no-cache",
                "Expires": "0"
            }
        )
    except Exception as e:
        logger.error(f"[EXPORT_ENDPOINT] Export error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel export workbook: {str(e)}")

