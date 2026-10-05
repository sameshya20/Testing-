import sys
import os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

AUTH_HEADERS = {"Authorization": "Bearer test-token-admin"}

def test_export_excel():
    # First generate test cases
    gen_payload = {
        "application_id": 1,
        "requirement_source": "USER_STORY",
        "requirement_identifier": "US-101",
        "application_url": None
    }
    gen_resp = client.post("/api/v1/generation/create", json=gen_payload, headers=AUTH_HEADERS)
    assert gen_resp.status_code == 200
    gen_data = gen_resp.json()

    export_payload = {
        "positive_test_cases": gen_data["positive_test_cases"],
        "negative_test_cases": gen_data["negative_test_cases"],
        "edge_test_cases": gen_data["edge_test_cases"],
        "application_name": "Enterprise Banking Portal",
        "requirement_ref": "US-101"
    }

    response = client.post("/api/v1/export/excel", json=export_payload, headers=AUTH_HEADERS)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Content-Disposition" in response.headers
    assert ".xlsx" in response.headers["Content-Disposition"]
    assert response.content[:4] == b'PK\x03\x04'
    assert len(response.content) > 1000

    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert set(wb.sheetnames) == {"Positive", "Negative", "Edge"}
    wb.close()

def test_export_multi_story_excel():
    gen_payload = {
        "application_id": 1,
        "requirement_source": "USER_STORY",
        "requirement_identifiers": ["US-101", "US-102"],
        "user_story_ids": ["US-101", "US-102"],
        "area_ids": [7]
    }
    gen_resp = client.post("/api/v1/generation/create", json=gen_payload, headers=AUTH_HEADERS)
    assert gen_resp.status_code == 200
    gen_data = gen_resp.json()

    export_payload = {
        "positive_test_cases": gen_data["positive_test_cases"],
        "negative_test_cases": gen_data["negative_test_cases"],
        "edge_test_cases": gen_data["edge_test_cases"],
        "application_name": "Enterprise Banking Portal",
        "requirement_ref": "US-101, US-102"
    }

    response = client.post("/api/v1/export/excel", json=export_payload, headers=AUTH_HEADERS)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert ".xlsx" in response.headers["Content-Disposition"]
    assert response.content[:4] == b'PK\x03\x04'
    
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert set(wb.sheetnames) == {"Positive", "Negative", "Edge"}
    pos_sheet = wb["Positive"]
    assert pos_sheet.max_row > 1
    wb.close()

