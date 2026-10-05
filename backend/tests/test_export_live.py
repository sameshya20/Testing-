import requests
import io
import openpyxl

def test_live_export():
    login_res = requests.post(
        'http://127.0.0.1:8000/api/v1/auth/login',
        json={'email': 'admin@enterprise.com', 'password': 'password123'}
    )
    assert login_res.status_code == 200, f"Login failed: {login_res.text}"
    token = login_res.json()['access_token']

    payload = {
        'positive_test_cases': [{
            'id': 'TC-POS-001',
            'requirement_ref': 'US-101',
            'user_story': 'US-101',
            'acceptance_criteria': 'AC-1',
            'application_page': 'Login Page',
            'application_link': 'http://localhost:5173/login',
            'category': 'Positive',
            'priority': 'High',
            'scenario': 'Valid Login with correct credentials',
            'description': 'User enters valid email and password',
            'preconditions': 'Registered user exists in DB',
            'test_data': 'email=admin@enterprise.com, password=password123',
            'test_steps': ['1. Navigate to /login', '2. Input credentials', '3. Click Sign In'],
            'expected_result': 'Dashboard renders with user profile'
        }],
        'negative_test_cases': [{
            'id': 'TC-NEG-001',
            'requirement_ref': 'US-101',
            'user_story': 'US-101',
            'acceptance_criteria': 'AC-1',
            'application_page': 'Login Page',
            'application_link': 'http://localhost:5173/login',
            'category': 'Negative',
            'priority': 'Medium',
            'scenario': 'Invalid password rejection',
            'description': 'User enters wrong password',
            'preconditions': 'Registered user exists in DB',
            'test_data': 'email=admin@enterprise.com, password=wrongpass',
            'test_steps': ['1. Navigate to /login', '2. Input wrong password', '3. Click Sign In'],
            'expected_result': 'Invalid credentials error banner appears'
        }],
        'edge_test_cases': [{
            'id': 'TC-EDG-001',
            'requirement_ref': 'US-101',
            'user_story': 'US-101',
            'acceptance_criteria': 'AC-1',
            'application_page': 'Login Page',
            'application_link': 'http://localhost:5173/login',
            'category': 'Edge',
            'priority': 'Low',
            'scenario': 'Blank field validation',
            'description': 'User submits empty form',
            'preconditions': 'None',
            'test_data': 'empty inputs',
            'test_steps': ['1. Navigate to /login', '2. Leave fields blank', '3. Click Sign In'],
            'expected_result': 'Required field validation highlights inputs'
        }],
        'application_name': 'Enterprise Application',
        'requirement_ref': 'US-101'
    }

    headers = {'Authorization': f'Bearer {token}'}
    res = requests.post('http://127.0.0.1:8000/api/v1/export/excel', json=payload, headers=headers)
    assert res.status_code == 200, f"Export failed: {res.text}"
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in res.headers.get('Content-Type', '')
    assert 'attachment' in res.headers.get('Content-Disposition', '')
    assert 'Test_Suite_US-101.xlsx' in res.headers.get('Content-Disposition', '')

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    print(f"Sheets: {wb.sheetnames}")
    assert set(wb.sheetnames) == {'Positive', 'Negative', 'Edge'}
    for name in ['Positive', 'Negative', 'Edge']:
        ws = wb[name]
        print(f"Sheet {name}: rows={ws.max_row}, cols={ws.max_column}")
        assert ws.max_row == 2  # header + 1 test case
        assert ws.max_column == 13  # all 13 standard columns
    print("ALL ASSERTIONS PASSED! Live export is 100% valid.")

if __name__ == '__main__':
    test_live_export()
