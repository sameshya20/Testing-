import io
import re
import os
import sys
sys.path.insert(0, os.path.abspath('backend'))
import xlsxwriter
import openpyxl
from typing import List, Any
from app.schemas.test_case import TestCase

class ExcelExporterServiceXLSXWriter:
    """Generates 100% Microsoft Excel-compliant multi-sheet workbooks using xlsxwriter."""

    def generate_workbook(
        self,
        positive_cases: List[TestCase],
        negative_cases: List[TestCase],
        edge_cases: List[TestCase],
        app_name: str,
        req_ref: str
    ) -> io.BytesIO:
        output = io.BytesIO()
        wb = xlsxwriter.Workbook(output, {'in_memory': True, 'strings_to_numbers': False})

        # Predefine styles / formats
        header_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 11,
            'bold': True,
            'font_color': '#FFFFFF',
            'bg_color': '#1E293B',
            'align': 'center',
            'valign': 'vcenter',
            'border': 1,
            'border_color': '#E2E8F0',
            'text_wrap': True
        })

        id_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 10,
            'bold': True,
            'font_color': '#0F172A',
            'align': 'center',
            'valign': 'top',
            'border': 1,
            'border_color': '#E2E8F0'
        })

        center_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 10,
            'font_color': '#0F172A',
            'align': 'center',
            'valign': 'top',
            'border': 1,
            'border_color': '#E2E8F0'
        })

        left_wrap_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 10,
            'font_color': '#0F172A',
            'align': 'left',
            'valign': 'top',
            'text_wrap': True,
            'border': 1,
            'border_color': '#E2E8F0'
        })

        p_high_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 10,
            'font_color': '#991B1B',
            'bg_color': '#FEE2E2',
            'align': 'center',
            'valign': 'top',
            'bold': True,
            'border': 1,
            'border_color': '#E2E8F0'
        })

        p_med_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 10,
            'font_color': '#92400E',
            'bg_color': '#FEF3C7',
            'align': 'center',
            'valign': 'top',
            'bold': True,
            'border': 1,
            'border_color': '#E2E8F0'
        })

        p_low_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 10,
            'font_color': '#075985',
            'bg_color': '#E0F2FE',
            'align': 'center',
            'valign': 'top',
            'border': 1,
            'border_color': '#E2E8F0'
        })

        empty_format = wb.add_format({
            'font_name': 'Calibri',
            'font_size': 10,
            'italic': True,
            'font_color': '#64748B'
        })

        headers = [
            "Test Case ID",
            "User Story",
            "Acceptance Criteria",
            "Application Page",
            "Application Link",
            "Test Scenario",
            "Test Description",
            "Preconditions",
            "Test Data",
            "Test Steps",
            "Expected Result",
            "Priority",
            "Category"
        ]

        sheets_data = [
            ("Positive", positive_cases),
            ("Negative", negative_cases),
            ("Edge", edge_cases)
        ]

        col_widths = [15, 16, 20, 20, 35, 35, 40, 35, 30, 45, 45, 14, 14]

        for sheet_idx, (sheet_name, cases) in enumerate(sheets_data):
            ws = wb.add_worksheet(sheet_name)
            ws.hide_gridlines(False) # show grid lines

            # Column Widths
            for col_idx, width in enumerate(col_widths):
                ws.set_column(col_idx, col_idx, width)

            # Header Row
            ws.set_row(0, 28)
            for col_idx, h in enumerate(headers):
                ws.write(0, col_idx, h, header_format)

            # Freeze header row
            ws.freeze_panes(1, 0)

            # Data rows
            if cases:
                for row_idx, tc in enumerate(cases, start=1):
                    if isinstance(tc.test_steps, list):
                        steps_text = "\n".join([str(s) for s in tc.test_steps if s is not None])
                    else:
                        steps_text = str(tc.test_steps or "")

                    user_story_val = getattr(tc, 'user_story', '') or ''
                    if not user_story_val:
                        user_story_val = tc.requirement_ref.split(' ')[0] if tc.requirement_ref else "N/A"

                    ac_val = getattr(tc, 'acceptance_criteria', '') or ''
                    if not ac_val:
                        parts = (tc.requirement_ref or '').split(' ')
                        ac_val = " ".join(parts[1:]) if len(parts) > 1 else (tc.requirement_ref or '')

                    app_page_val = getattr(tc, 'application_page', '') or "N/A"
                    app_link_val = getattr(tc, 'application_link', '') or "N/A"

                    # Calculate dynamic row height
                    step_lines = len(steps_text.split('\n'))
                    desc_lines = max(1, len(str(tc.description or '')) // 35)
                    row_h = max(45, max(step_lines, desc_lines) * 18)
                    ws.set_row(row_idx, min(400, row_h))

                    # Row values
                    row_values = [
                        (tc.id, id_format),
                        (user_story_val, center_format),
                        (ac_val, center_format),
                        (app_page_val, center_format),
                        (app_link_val, left_wrap_format),
                        (tc.scenario, left_wrap_format),
                        (tc.description, left_wrap_format),
                        (tc.preconditions, left_wrap_format),
                        (tc.test_data, left_wrap_format),
                        (steps_text, left_wrap_format),
                        (tc.expected_result, left_wrap_format),
                        (tc.priority, None), # Handled below
                        (tc.category, center_format)
                    ]

                    for col_num, (val, fmt) in enumerate(row_values):
                        if col_num == 11: # Priority
                            p_str = str(val or 'Medium').upper()
                            if 'HIGH' in p_str:
                                p_fmt = p_high_format
                            elif 'MED' in p_str:
                                p_fmt = p_med_format
                            else:
                                p_fmt = p_low_format
                            ws.write(row_idx, col_num, str(val or 'Medium'), p_fmt)
                        else:
                            ws.write(row_idx, col_num, str(val or ''), fmt)

                # Auto-filter
                ws.autofilter(0, 0, len(cases), len(headers) - 1)
            else:
                ws.set_row(1, 24)
                ws.write(1, 0, f"No {sheet_name.lower()} test cases generated for this requirement.", empty_format)

        # Set first sheet active
        wb.worksheets()[0].activate()
        wb.close()
        output.seek(0)
        return output

# Let's test this!
service = ExcelExporterServiceXLSXWriter()
# Load sample test cases from previous inspection
import json
import urllib.request

req = urllib.request.Request(
    "http://127.0.0.1:8000/api/v1/generation/create",
    data=json.dumps({
        "application_id": 1,
        "requirement_source": "USER_STORY",
        "requirement_identifiers": ["US-101", "US-103"],
        "user_story_ids": ["US-101", "US-103"],
        "area_ids": [7]
    }).encode("utf-8"),
    headers={"Content-Type": "application/json", "Authorization": "Bearer test-token-admin"}
)
with urllib.request.urlopen(req) as resp:
    data = json.loads(resp.read().decode())
    from app.schemas.test_case import TestCase
    pos = [TestCase(**c) for c in data['positive_test_cases']]
    neg = [TestCase(**c) for c in data['negative_test_cases']]
    edge = [TestCase(**c) for c in data['edge_test_cases']]

res_stream = service.generate_workbook(pos, neg, edge, "Test App", "US-101, US-103")
content = res_stream.getvalue()

with open("test_perfect.xlsx", "wb") as f:
    f.write(content)
print(f"Generated test_perfect.xlsx, size: {len(content)} bytes")

# Test with openpyxl
wb_check = openpyxl.load_workbook("test_perfect.xlsx")
print("openpyxl sheet names:", wb_check.sheetnames)
assert wb_check.sheetnames == ["Positive", "Negative", "Edge"]
print("Positive rows:", wb_check["Positive"].max_row)
wb_check.close()
print("SUCCESS: test_perfect.xlsx is valid!")
