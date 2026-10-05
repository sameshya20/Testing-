import os
import subprocess
import requests
import io
import openpyxl

BASE_URL = "http://127.0.0.1:8000/api/v1"

# 1. Login
login_res = requests.post(
    f"{BASE_URL}/auth/login",
    json={"email": "admin@enterprise.com", "password": "password123"}
)
assert login_res.status_code == 200
token = login_res.json()["access_token"]
headers = {"Authorization": f"Bearer {token}"}

scenarios = [
    {
        "name": "Single User Story US-101",
        "req_payload": {
            "application_id": 1,
            "requirement_source": "USER_STORY",
            "requirement_identifiers": ["US-101"],
            "user_story_ids": ["US-101"],
            "area_ids": [7]
        },
        "req_ref": "US-101"
    },
    {
        "name": "Multiple User Stories US-101 and US-103",
        "req_payload": {
            "application_id": 1,
            "requirement_source": "USER_STORY",
            "requirement_identifiers": ["US-101", "US-103"],
            "user_story_ids": ["US-101", "US-103"],
            "area_ids": [7]
        },
        "req_ref": "US-101, US-103"
    },
    {
        "name": "Multiple User Stories US-101, US-102, US-103",
        "req_payload": {
            "application_id": 1,
            "requirement_source": "USER_STORY",
            "requirement_identifiers": ["US-101", "US-102", "US-103"],
            "user_story_ids": ["US-101", "US-102", "US-103"],
            "area_ids": [7]
        },
        "req_ref": "US-101, US-102, US-103"
    }
]

for idx, sc in enumerate(scenarios, 1):
    print(f"\n=======================================================")
    print(f"RUNNING SCENARIO {idx}: {sc['name']}")
    print(f"=======================================================")

    # Generate
    gen_res = requests.post(f"{BASE_URL}/generation/create", json=sc["req_payload"], headers=headers)
    assert gen_res.status_code == 200
    gen_data = gen_res.json()
    pos = gen_data["positive_test_cases"]
    neg = gen_data["negative_test_cases"]
    edge = gen_data["edge_test_cases"]
    print(f"[+] Generated: {len(pos)} Positive, {len(neg)} Negative, {len(edge)} Edge")

    # Export
    export_payload = {
        "positive_test_cases": pos,
        "negative_test_cases": neg,
        "edge_test_cases": edge,
        "application_name": "Enterprise Banking",
        "requirement_ref": sc["req_ref"]
    }
    exp_res = requests.post(f"{BASE_URL}/export/excel", json=export_payload, headers=headers)
    assert exp_res.status_code == 200
    content_type = exp_res.headers.get("Content-Type", "")
    disposition = exp_res.headers.get("Content-Disposition", "")
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in content_type
    assert ".xlsx" in disposition
    assert not ".xls\"" in disposition
    print(f"[+] Response Content-Type: {content_type}")
    print(f"[+] Response Content-Disposition: {disposition}")

    content = exp_res.content
    assert content[:4] == b'PK\x03\x04', "Missing PK signature"
    print(f"[+] Binary size: {len(content)} bytes, signature: {content[:4]!r}")

    # Save to file
    out_file = os.path.abspath(f"scratch/scenario_{idx}_export.xlsx")
    with open(out_file, "wb") as f:
        f.write(content)

    # Verify openpyxl
    wb = openpyxl.load_workbook(out_file)
    assert set(wb.sheetnames) == {"Positive", "Negative", "Edge"}
    for sh in ["Positive", "Negative", "Edge"]:
        assert wb[sh].max_row > 1
        assert wb[sh].max_column == 13
    wb.close()
    print(f"[+] openpyxl verified sheets and columns successfully.")

    # Verify Microsoft Excel COM
    ps_cmd = f"""
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $true
    try {{
        $wb = $excel.Workbooks.Open('{out_file}')
        Write-Host "Excel Opened: $($wb.Name), Sheets: $($wb.Sheets.Count)"
        foreach ($s in $wb.Sheets) {{
            Write-Host "  $($s.Name): $($s.UsedRange.Rows.Count) rows"
        }}
        $wb.Close($false)
        Write-Host "SUCCESS"
    }} catch {{
        Write-Error "ERROR: $_"
        exit 1
    }} finally {{
        $excel.Quit()
        [System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) | Out-Null
    }}
    """
    res = subprocess.run(["powershell", "-NoProfile", "-Command", ps_cmd], capture_output=True, text=True)
    print(res.stdout)
    assert res.returncode == 0, f"Excel failed for scenario {idx}: {res.stderr}"
    assert "SUCCESS" in res.stdout

print("\n>>> ALL MULTIPLE DOWNLOAD SCENARIOS PASSED WITH 100% SUCCESS IN MICROSOFT EXCEL!")
